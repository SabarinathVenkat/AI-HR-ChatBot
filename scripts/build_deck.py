"""Build the demo deck (PowerPoint) from the project's real screenshots, diagrams and evaluation numbers.

    pip install python-pptx            # only this script needs it, so it is not in requirements
    python scripts/build_deck.py       # writes docs/deck/HR_Chat_Agent_Demo_Deck.pptx

The slide text lives in this file: edit it here and rebuild. Numbers come from docs/EVALUATION.md and
the screenshots from scripts/ui_screenshots.py, so re-run that script after any UI change.
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
PNG = ROOT / "docs" / "diagrams" / "architecture" / "png"
OUT = ROOT / "docs" / "deck" / "HR_Chat_Agent_Demo_Deck.pptx"

NAVY, NAVY2 = "0A1730", "0F2347"
BLUE, BLUE_SOFT = "2563EB", "E8F0FF"
TEAL, TEAL_SOFT = "0D9488", "DFF5F1"
PURPLE, PURPLE_SOFT = "7C3AED", "EDE7FD"
AMBER, AMBER_SOFT = "B45309", "FFF4D6"
RED, RED_SOFT = "B3261E", "FDE8E6"
INK, MUTED, LINE, BG, WHITE = "0F1B2D", "556278", "D5DCE8", "F2F5FA", "FFFFFF"
LIGHT = "C4D2EA"
FONT = "Calibri"

KINDS = {  # fill, outline, text colour
    "plain": (WHITE, "9AA8C0", INK),
    "focal": (BLUE, BLUE, WHITE),
    "llm": (BLUE_SOFT, BLUE, INK),
    "store": (BG, "7A8399", INK),
    "ext": (BG, "9AA8C0", INK),
    "security": (AMBER_SOFT, AMBER, INK),
    "plan": (PURPLE_SOFT, PURPLE, INK),
    "ok": (TEAL_SOFT, TEAL, INK),
    "bad": (RED_SOFT, RED, INK),
    "dark": (NAVY2, NAVY2, WHITE),
}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
count = {"n": 0}


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------- text ----------
def _bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.27))))
    pPr.set("indent", str(-int(Inches(0.27))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def fill_paras(tf, paras, size, color, bold, align, space_after):
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        spec = para if isinstance(para, dict) else {"runs": [para]}
        runs = spec["runs"]
        p.alignment = spec.get("align", align)
        p.space_after = Pt(spec.get("space_after", space_after))
        if spec.get("bullet"):
            _bullet(p)
        for run in runs:
            text, opts = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(opts.get("size", spec.get("size", size)))
            r.font.bold = opts.get("bold", spec.get("bold", bold))
            r.font.italic = opts.get("italic", False)
            r.font.color.rgb = rgb(opts.get("color", spec.get("color", color)))


def bullets(items, **extra):
    return [{"runs": [item] if isinstance(item, (str, tuple)) else item, "bullet": True, **extra} for item in items]


def add_text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    fill_paras(tf, paras, size, color, bold, align, space_after)
    return tb


# ---------- shapes ----------
def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, paras=None, size=14, color=INK, bold=False,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1, line_w=1.25, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, dash=False, space_after=2, margin=0.1):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = anchor
    if paras:
        fill_paras(tf, paras, size, color, bold, align, space_after)
    return shp


def node(slide, x, y, w, h, title, sub=None, kind="plain", size=15, sub_size=11.5, dash=False):
    fill, line, text = KINDS[kind]
    sub_color = LIGHT if kind in ("focal", "dark") else MUTED
    paras = [{"runs": [title], "bold": True, "size": size, "color": text, "space_after": 1}]
    if sub:
        paras += [{"runs": [line_], "size": sub_size, "color": sub_color, "space_after": 0} for line_ in
                  ([sub] if isinstance(sub, str) else sub)]
    return add_box(slide, x, y, w, h, fill=fill, line=line, paras=paras, dash=dash)


def card(slide, x, y, w, h, title, items, accent=BLUE, size=15, title_size=17, fill=WHITE):
    add_box(slide, x, y, w, h, fill=fill, line=LINE, radius=0.04)
    add_box(slide, x, y, 0.09, h, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, x + 0.25, y + 0.15, w - 0.4, 0.5, [title], size=title_size, bold=True, color=accent)
    add_text(slide, x + 0.25, y + 0.7, w - 0.45, h - 0.8, bullets(items, space_after=7), size=size)


def arrow(slide, pts, color=MUTED, width=2.0, dash=False, head=True):
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        if dash:
            c.line.dash_style = MSO_LINE.DASH
        if head and i == len(pts) - 2:
            ln = c.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")


def label(slide, x, y, w, text, size=12, color=MUTED, align=PP_ALIGN.CENTER, bold=False):
    return add_text(slide, x, y, w, 0.35, [text], size=size, color=color, align=align, bold=bold)


def picture(slide, path, x, y, w, h, border=True):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    left, top = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(pw), Inches(ph))
    if border:
        pic.line.color.rgb = rgb(LINE)
        pic.line.width = Pt(0.75)
    return left, top, pw, ph


# ---------- slide scaffolding ----------
def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    count["n"] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def content_slide(title, subtitle=None, notes=""):
    s = new_slide()
    add_box(s, 0.6, 0.4, 0.7, 0.07, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 0.55, 0.5, 12.2, 0.7, [title], size=30, color=NAVY, bold=True)
    if subtitle:
        add_text(s, 0.55, 1.15, 12.2, 0.45, [subtitle], size=16, color=MUTED)
    add_text(s, 0.55, 7.02, 9, 0.3, ["HR Chat Agent  ·  sample data and sample policies are illustrative"], size=10.5, color=MUTED)
    add_text(s, 12.0, 7.02, 0.75, 0.3, [str(count["n"])], size=10.5, color=MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


# =====================================================================
# 1  Title
# =====================================================================
s = new_slide(NAVY)
add_box(s, 9.9, 1.85, 2.6, 2.6, fill=BLUE, line=None, paras=["HR"], size=96, color=WHITE, bold=True, radius=0.2)
add_text(s, 0.8, 1.65, 8.8, 1.3, ["HR Chat Agent"], size=60, bold=True, color=WHITE)
add_text(s, 0.8, 3.0, 8.8, 1.1, ["An HR assistant you can trust with real numbers"], size=28, color=LIGHT)
add_text(s, 0.8, 4.4, 8.8, 1.0,
         ["Answers from policy  ·  reads your own records  ·  never submits anything without your confirmation"],
         size=18, color=LIGHT)
add_text(s, 0.8, 6.2, 9, 0.5, ["AI Master evaluation  ·  Sabari"], size=16, color="8FA6C9")
add_text(s, 0.8, 6.7, 11, 0.4, ["Sample data and sample policies, written for this demonstration"], size=12, color="8FA6C9")
s.notes_slide.notes_text_frame.text = (
    "Open with the one-line promise: an HR assistant you can trust with real numbers. Everything in this deck "
    "serves that. The assistant answers from policy, reads the employee's own records, and cannot submit anything "
    "without the employee confirming. Mention that the data and policies are samples written for the demonstration. "
    "About 15 minutes: a live demo first, then how it works, how we know it works, and where it falls short."
)

# =====================================================================
# 2  Problem and brief
# =====================================================================
s = content_slide(
    "The problem we set out to solve",
    notes=(
        "HR questions repeat constantly, and the answers sit in policy documents and HR systems. Leave questions are "
        "the hard kind: they depend on a rule and on the person's own data. A chatbot that guesses a number loses trust "
        "immediately. The brief asked for authenticated employees, answers from the policy documents, tools and a "
        "database for the dynamic questions, and documented architecture. Our answer is one idea: let the model talk, "
        "and let code decide facts and permissions."
    ),
)
card(s, 0.6, 1.8, 5.9, 3.7, "What employees face", [
    "Policy answers are scattered across documents",
    "Leave balances and eligibility need rules plus personal data, not guesses",
    "One wrong number in an HR answer costs trust, and can cost money",
], accent=TEAL, size=18)
card(s, 6.8, 1.8, 5.9, 3.7, "What the brief asked for", [
    "Employees sign in, and every answer is for the right person",
    "Answers come from the HR policy documents, with their source",
    "Tools and a database answer the dynamic questions: leave balance, eligibility, calculations",
    "Architecture that is documented and can be evaluated",
], accent=BLUE, size=18)
add_box(s, 0.6, 5.8, 12.1, 0.95, fill=NAVY, line=None,
        paras=["Our answer: the model talks, and code decides facts and permissions."], size=22, color=WHITE, bold=True)

# =====================================================================
# 3  What it does
# =====================================================================
s = content_slide(
    "What it does",
    notes=(
        "Walk down the list. Policy answers carry clause citations, and each citation opens the policy text. Personal "
        "data comes from the employee's own records. Day counts and eligibility are computed by tested rule code, not "
        "by the model. Applying for leave produces a draft the employee confirms. Managers get approvals and team "
        "availability for their own direct reports. Planning compares leave options across holidays. The screenshot is "
        "the real application after sign-in as Ravi, an employee."
    ),
)
rows = [
    ("Policy answers", "with clause citations that open the policy text"),
    ("Your leave data", "balances, requests and holidays from your own records"),
    ("Day counts and eligibility", "from tested rule code, not from the model"),
    ("Apply for leave", "the assistant prepares a draft, you confirm it"),
    ("Manager tools", "pending approvals and team availability, for direct reports only"),
    ("Planning", "compare leave options across holidays"),
]
add_text(s, 0.6, 1.8, 5.9, 5.0,
         [{"runs": [(lead, {"bold": True, "color": NAVY}), "  " + rest], "space_after": 14} for lead, rest in rows],
         size=18)
picture(s, SHOTS / "02-home.png", 6.8, 1.8, 5.95, 4.95)

# =====================================================================
# 4  Demo: policy
# =====================================================================
s = content_slide(
    "Demo 1: a policy question, answered with its source",
    notes=(
        "Live demo: sign in as Ravi and ask what the notice period is after resigning. Point at the answer: each rule "
        "has a clause chip like POL-05 section 3.1, and clicking it opens the policy text. 'How I got this' shows the "
        "tool call behind the answer, and the tag shows the detected intent and the time. Explain that the text appears "
        "only after the verifier has checked it, which is why the answer arrives as a whole and not word by word. "
        "Note that the assistant says the policies are illustrative samples."
    ),
)
left, top, pw, ph = picture(s, SHOTS / "03-policy-answer.png", 0.6, 1.75, 7.9, 4.95)
add_text(s, 0.6, 6.68, 7.9, 0.3, ["Asked: “What is the notice period after I resign?”  Answered in 5.6 s."], size=12, color=MUTED)
callouts = [
    ("Real source", "Each rule carries a clause chip, like POL-05 §3.1. Clicking it opens the policy text."),
    ("Visible steps", "“How I got this” lists the tool calls behind the answer."),
    ("Verified first", "The text is sent only after numbers and citations have been checked."),
    ("Honest", "The assistant says the sample policies are illustrative."),
]
for i, (head, body) in enumerate(callouts):
    y = 1.8 + i * 1.28
    add_box(s, 8.85, y + 0.05, 0.5, 0.5, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL, paras=[str(i + 1)], size=16, color=WHITE, bold=True)
    add_text(s, 9.5, y, 3.3, 1.2, [{"runs": [head], "bold": True, "size": 16, "color": NAVY, "space_after": 2}, {"runs": [body], "size": 13.5}])

# =====================================================================
# 5  Demo: leave
# =====================================================================
s = content_slide(
    "Demo 2: applying for leave. The model proposes, you confirm.",
    notes=(
        "Live demo: ask to take annual leave from 2 to 6 November. The assistant checks every rule and stores a draft; "
        "it shows a confirmation card with the leave type, dates and working days, and says nothing has been submitted. "
        "Click Confirm. The request appears as pending in the side panel and the manager is asked to approve. The key "
        "design point: the model never sees the draft id and has no tool that submits. Only the Confirm click, or a "
        "bare 'yes' handled by code, creates the request, and confirmation re-checks eligibility."
    ),
)
picture(s, SHOTS / "05-confirm-card.png", 0.6, 1.75, 6.0, 3.95)
picture(s, SHOTS / "06-submitted.png", 6.75, 1.75, 6.0, 3.95)
label(s, 0.6, 5.75, 6.0, "1   Draft prepared. Nothing is submitted yet.", size=14, color=NAVY, bold=True)
label(s, 6.75, 5.75, 6.0, "2   After Confirm: the request is pending and the manager is asked.", size=14, color=NAVY, bold=True)
add_box(s, 0.6, 6.25, 12.15, 0.65, fill=AMBER_SOFT, line=None, size=14, color=INK,
        paras=["The model can prepare a draft but never sees its id. Only the Confirm click (or a bare “yes”, handled by code) "
               "creates the request, and the rules are checked again at that moment."])

# =====================================================================
# 6  Managers and phone
# =====================================================================
s = content_slide(
    "Managers, and the phone",
    notes=(
        "Switch to Arjun, a manager. His panel shows requests waiting for his approval, and only from his direct "
        "reports: the filter is inside the database query, so a bug in the agent cannot widen access. Approving or "
        "rejecting is one click. The phone screenshot is the same app at 390 pixels wide; a real-browser test checks "
        "that there is no sideways scrolling."
    ),
)
picture(s, SHOTS / "07-manager.png", 0.6, 1.75, 6.9, 4.9)
picture(s, SHOTS / "08-phone.png", 7.75, 1.75, 2.3, 4.9)
add_text(s, 10.3, 1.9, 2.5, 4.8, bullets([
    "Arjun sees requests from his direct reports only. The filter is in the SQL.",
    "Approve or reject in one click.",
    "The same page works at 390 px wide with no sideways scroll, checked in a real browser.",
], space_after=14), size=14.5)

# =====================================================================
# 7  Central idea
# =====================================================================
s = content_slide(
    "The central idea: the model talks, code decides",
    "Five questions, and who answers each one",
    notes=(
        "This is the design principle behind everything else. Which employee is asking comes from the signed token, "
        "never from the model. Every number, date and eligibility result comes from tested Python reached through "
        "tools. Policy wording comes from retrieved sections cited by clause. Whether an answer may be shown is "
        "decided by a deterministic verifier. Whether leave is submitted is decided by the employee through an "
        "endpoint the model cannot call. The model is good at language; it is never the authority on facts or permissions."
    ),
)
rules = [
    ("Which employee is asking", "The signed login token. Never the model."),
    ("Every number, date and eligibility result", "Tested Python rules, reached through tools."),
    ("Policy wording", "Retrieved policy sections, cited by clause id."),
    ("Whether an answer may be shown", "A deterministic verifier, after every draft."),
    ("Whether leave is submitted", "The employee, through an endpoint the model cannot call."),
]
for i, (what, who) in enumerate(rules):
    y = 1.85 + i * 1.0
    add_box(s, 0.6, y, 4.6, 0.88, fill=NAVY, line=None, paras=[what], size=17, color=WHITE, bold=True, align=PP_ALIGN.LEFT, margin=0.2)
    arrow(s, [(5.25, y + 0.44), (5.65, y + 0.44)], color=BLUE, width=2.5)
    add_box(s, 5.7, y, 7.05, 0.88, fill=BG, line=LINE, paras=[who], size=17, align=PP_ALIGN.LEFT, margin=0.2)

# =====================================================================
# 8  Architecture
# =====================================================================
s = content_slide(
    "Architecture at a glance",
    "Custom web app  ·  FastAPI  ·  LangGraph  ·  Deep Agents  ·  Claude  ·  one SQLite file",
    notes=(
        "Read left to right. Users reach a custom web app, which talks to a FastAPI gateway that checks the login "
        "token, rate limits and attaches the verified identity. The LangGraph agent is the core: it calls Claude, "
        "uses ten scoped tools built on tested leave rules, searches policies, and for planning questions hands off "
        "to a read-only Deep Agents path. Everything lands in one SQLite file: HR data, leave drafts, the policy "
        "index, the audit log and the chat history. Choosing SQLite keeps setup to zero for reviewers; the trade-off "
        "is a single writer, and moving to Postgres is a documented migration."
    ),
)
node(s, 0.6, 3.65, 1.95, 0.95, "Web UI", ["employee · manager · HR"], "ext")
node(s, 3.05, 3.65, 2.1, 0.95, "FastAPI gateway", ["JWT · rate limit · SSE"], "plain")
node(s, 5.6, 3.3, 2.3, 1.65, "LangGraph agent", ["router · tool loop", "verifier"], "focal", size=17)
node(s, 5.6, 1.8, 2.3, 0.95, "Claude API", ["Sonnet answers · Haiku routes"], "ext")
node(s, 8.65, 1.8, 2.2, 0.95, "Deep Agents", ["planning path · read-only"], "plan")
node(s, 8.65, 3.65, 2.2, 0.95, "Tools + rules", ["10 scoped tools · R1-R10"], "plain")
node(s, 8.65, 5.45, 2.2, 0.95, "Policy search", ["FTS5 + vectors · cited"], "plain")
node(s, 11.25, 3.3, 1.5, 1.65, "SQLite", ["one file"], "store")
arrow(s, [(2.55, 4.12), (3.05, 4.12)])
arrow(s, [(5.15, 4.12), (5.6, 4.12)])
arrow(s, [(6.75, 3.3), (6.75, 2.75)], color=BLUE)
arrow(s, [(7.9, 3.6), (8.27, 3.6), (8.27, 2.27), (8.65, 2.27)])
arrow(s, [(7.9, 4.12), (8.65, 4.12)])
arrow(s, [(7.9, 4.65), (8.27, 4.65), (8.27, 5.92), (8.65, 5.92)])
arrow(s, [(9.75, 2.75), (9.75, 3.65)], dash=True, color=PURPLE)
label(s, 9.85, 2.98, 1.6, "read-only", size=11, color=PURPLE, align=PP_ALIGN.LEFT)
arrow(s, [(10.85, 4.12), (11.25, 4.12)])
arrow(s, [(10.85, 5.92), (12.0, 5.92), (12.0, 4.95)])
add_text(s, 0.6, 5.35, 7.0, 1.4,
         [{"runs": [("One SQLite file", {"bold": True, "color": NAVY}), " holds HR data, leave drafts, the policy index "
           "(FTS5 + sqlite-vec), the audit log and the chat history."], "size": 15}],
         size=15)

# =====================================================================
# 9  Two routes
# =====================================================================
s = content_slide(
    "Two routes: everyday questions and complex plans",
    "Plain code picks the route. The model never chooses its own safety checks.",
    notes=(
        "Every message passes the input guardrail, then a small fast model, Haiku, labels the intent with a confidence. "
        "A plain Python function, not the model, chooses the route. Most questions take the everyday route: Sonnet picks "
        "from the tools its role allows, up to five calls, and drafts a cited answer. Small routes handle clarifying "
        "questions, out-of-scope requests, small talk and sensitive topics, where code even chooses the tools. "
        "Multi-step planning goes to Deep Agents. All drafted answers pass the same verifier. A failed draft is "
        "rewritten up to twice and then replaced by a plain answer built from the tool output. If planning fails, the "
        "everyday tool loop answers instead."
    ),
)
node(s, 0.5, 3.55, 1.45, 0.95, "Message", ["from /chat"], "ext")
node(s, 2.2, 3.55, 1.6, 0.95, "Guardrail", ["length · injection"], "plain")
node(s, 4.05, 3.55, 1.7, 0.95, "Router", ["Haiku · intent"], "llm")
arrow(s, [(1.95, 4.02), (2.2, 4.02)])
arrow(s, [(3.8, 4.02), (4.05, 4.02)])
arrow(s, [(5.75, 4.02), (6.2, 4.02)], head=False)
arrow(s, [(6.2, 2.25), (6.2, 5.45)], head=False)
arrow(s, [(6.2, 2.25), (6.65, 2.25)])
arrow(s, [(6.2, 4.02), (6.65, 4.02)])
arrow(s, [(6.2, 5.45), (6.65, 5.45)])
label(s, 3.9, 4.6, 2.3, "Code, not the model, picks the route", size=12, color=BLUE)
node(s, 6.65, 1.8, 3.1, 0.95, "Everyday: agent + tools", ["Sonnet picks tools · max 5 calls"], "llm")
node(s, 6.65, 3.5, 3.1, 1.05, "Small routes", ["clarify · decline · small talk", "sensitive (fixed tools)"], "plain")
node(s, 6.65, 5.0, 3.1, 0.95, "Complex: Deep Agents", ["planner + 2 sub-agents · read-only"], "plan")
arrow(s, [(9.75, 2.27), (10.2, 2.27)])
arrow(s, [(9.75, 4.02), (10.2, 4.02)])
arrow(s, [(9.75, 5.47), (10.2, 5.47)])
node(s, 10.2, 1.8, 1.45, 4.15, "Verify", ["numbers", "dates", "citations", "names"], "security", size=17)
arrow(s, [(11.65, 3.87), (11.9, 3.87)])
node(s, 11.9, 3.4, 0.95, 0.95, "Answer", None, "ok", size=12.5)
add_text(s, 6.65, 6.0, 3.2, 0.7, ["If planning fails, the everyday tool loop answers."], size=12, color=MUTED)
add_text(s, 9.9, 6.0, 3.0, 0.8, ["Failed draft: rewrite up to 2 times, then a plain fallback built from tool output."], size=12, color=MUTED)

# =====================================================================
# 10  Why this split
# =====================================================================
s = content_slide(
    "Why LangGraph for the spine, Deep Agents for planning",
    notes=(
        "LangGraph gives us a fixed spine. Safety steps are graph nodes the model cannot skip. Routing is plain code, "
        "so behaviour is predictable and testable. Conversation memory and live progress streaming are built in, and "
        "the tests run the whole graph with a scripted fake model. Deep Agents earns its place only where planning and "
        "delegation help: multi-step questions. Sub-agents get clean contexts and the faster model. We fence it hard, "
        "because its defaults include file and shell tools: read-only tools, 12 tool calls, 3 sub-agent runs and 22 "
        "seconds. Be upfront about the cost: the planning path is slow, median 32 seconds, which is why simple "
        "questions never take it."
    ),
)
card(s, 0.6, 1.8, 6.0, 4.3, "LangGraph: the fixed spine", [
    "Safety steps are nodes the model cannot skip: guardrail, verify, rewrite, fallback",
    "Routing is plain code, so behaviour is predictable and testable",
    "Conversation memory and live progress streaming come built in",
    "Tests run the whole graph with a scripted fake model",
], accent=BLUE, size=16)
card(s, 6.8, 1.8, 5.95, 4.3, "Deep Agents: only for planning", [
    "Planning and delegation suit multi-step questions",
    "Sub-agents get clean contexts and a faster model",
    "Fenced in: read-only tools, 12 tool calls, 3 sub-agent runs, 22 seconds",
    "If it fails, the normal tool loop answers instead",
], accent=PURPLE, size=16)
add_box(s, 0.6, 6.3, 12.15, 0.6, fill=AMBER_SOFT, line=None, size=14.5,
        paras=["Honest cost: the planning path is slow (median 32 s). That is why simple questions never take it."])

# =====================================================================
# 11  Policy RAG
# =====================================================================
s = content_slide(
    "Policy answers: hybrid search, and a gate that says “not found”",
    "Everything lives in one SQLite file. Local embeddings, so no second API key.",
    notes=(
        "Offline, the eight policy files are split into one chunk per numbered section, 53 sections in total, each "
        "carrying its policy id, version and clause ids. Each chunk is embedded locally with bge-small and stored in "
        "SQLite next to a keyword index. Online, a question searches keywords and meaning in parallel, takes the top "
        "50 of each, and merges them by reciprocal rank fusion into the top five. There is no reranker; with 53 "
        "sections it was not needed. The confidence gate has three tiers: confident at 0.655 or higher, a flagged "
        "weak match between 0.50 and 0.655, and nothing below 0.50, where the assistant says it could not find the "
        "answer and gives the HR contact. The weak tier exists because evaluation showed rephrased questions score "
        "lower even when the policy answers them."
    ),
)
label(s, 0.6, 1.62, 3.5, "INDEXING, ONCE AT STARTUP", size=11.5, color=MUTED, align=PP_ALIGN.LEFT, bold=True)
xs = [0.6, 3.8, 7.0, 10.2]
top_row = [("Policy files", "8 markdown documents", "ext"), ("Chunker", "one chunk per section", "plain"),
           ("Embeddings", "bge-small · local", "plain"), ("SQLite index", "text · keywords · vectors", "store")]
for x, (t, sub_, kind) in zip(xs, top_row, strict=True):
    node(s, x, 1.95, 2.45, 1.0, t, [sub_], kind)
for i in range(3):
    arrow(s, [(xs[i] + 2.45, 2.45), (xs[i + 1], 2.45)])
label(s, 0.6, 3.35, 3.5, "EACH QUESTION", size=11.5, color=MUTED, align=PP_ALIGN.LEFT, bold=True)
bottom_row = [("Question", "from the search tool", "ext"), ("Hybrid search", "keywords top 50 + vectors top 50", "plain"),
              ("Fuse", "reciprocal rank · top 5", "plain"), ("Confidence tiers", "best vector similarity", "security")]
for x, (t, sub_, kind) in zip(xs, bottom_row, strict=True):
    node(s, x, 3.7, 2.45, 1.0, t, [sub_], kind)
for i in range(3):
    arrow(s, [(xs[i] + 2.45, 4.2), (xs[i + 1], 4.2)])
arrow(s, [(11.42, 2.95), (11.42, 3.32), (5.0, 3.32), (5.0, 3.7)], color=BLUE, dash=True)
tiers = [
    ("0.655 or higher: confident", "The top 5 sections go to Claude, which cites them.", TEAL, TEAL_SOFT),
    ("0.50 to 0.655: weak match", "Up to 3 sections, flagged. Used only if one directly answers.", AMBER, AMBER_SOFT),
    ("Below 0.50: not found", "Say so and give the HR contact. No guessing.", RED, RED_SOFT),
]
for i, (head, body, accent, soft) in enumerate(tiers):
    x = 0.6 + i * 4.1
    add_box(s, x, 5.1, 3.95, 1.45, fill=soft, line=accent, radius=0.06,
            paras=[{"runs": [head], "bold": True, "size": 16, "color": accent, "space_after": 4}, {"runs": [body], "size": 14}],
            align=PP_ALIGN.LEFT, margin=0.2)
add_text(s, 0.6, 6.62, 12.1, 0.35, ["No reranker: with 53 sections it was not needed. Citations are plain text, and the verifier checks each one was retrieved."], size=12, color=MUTED)

# =====================================================================
# 12  Security
# =====================================================================
s = content_slide(
    "Security: the model is treated as untrusted",
    "It can propose a tool and arguments. The server decides the rest.",
    notes=(
        "Follow the chain. The model proposes a tool name and arguments, such as dates and a leave type, and nothing "
        "else. The server adds the employee id and role from the signed token. The tool call is checked: is the tool "
        "allowed for this role, and are the arguments valid? Unknown arguments, like an employee id, are rejected and "
        "audited. Tools call pure, tested domain rules, and repositories whose SQL is filtered by employee or "
        "manager id, so a manager can only ever see direct reports because the filter is in the query itself. Every "
        "call, including refused ones, is written to the audit log. Below are the other protections: argon2 password "
        "hashes, 60-minute tokens, lockout after five failures, a per-employee rate limit, and a strict "
        "content-security policy."
    ),
)
cols = [0.6, 3.1, 5.6, 8.1, 10.6]
chain = [
    ("Model proposal", ["tool name and", "arguments only"], "ext"),
    ("Server adds identity", ["id and role from", "the signed token"], "security"),
    ("Tool call checked", ["role allowed?", "unknown args rejected"], "plain"),
    ("Scoped repository", ["SQL filtered by", "employee or manager"], "plain"),
    ("SQLite", ["HR data, drafts"], "store"),
]
for x, (t, sub_, kind) in zip(cols, chain, strict=True):
    node(s, x, 1.85, 2.0, 1.15, t, sub_, kind, size=14.5)
for i in range(4):
    arrow(s, [(cols[i] + 2.0, 2.42), (cols[i + 1], 2.42)])
node(s, 5.6, 3.5, 2.0, 0.95, "Domain rules", ["pure, tested code"], "ok", size=14.5)
node(s, 8.1, 3.5, 2.0, 0.95, "audit_log", ["every call, even refused"], "store", size=14.5)
arrow(s, [(6.6, 3.0), (6.6, 3.5)])
arrow(s, [(9.1, 3.0), (9.1, 3.5)], dash=True)
protections = [
    "Identity never comes from the model: a call that includes an employee id is rejected and audited.",
    "Managers see direct reports only. The filter is inside the SQL query, not in the agent.",
    "The model can prepare a leave draft but cannot submit it.",
    "argon2 passwords, 60-minute tokens, lockout after 5 failures, 30 messages a minute, strict content-security policy.",
]
for i, text in enumerate(protections):
    x = 0.6 + (i % 2) * 6.2
    y = 4.85 + (i // 2) * 1.05
    add_box(s, x, y, 5.95, 0.92, fill=BG, line=LINE, paras=[text], size=14, align=PP_ALIGN.LEFT, margin=0.18)

# =====================================================================
# 13  Verifier
# =====================================================================
s = content_slide(
    "Never trust the model: the verifier",
    "A deterministic check after every draft. It is code, not another model.",
    notes=(
        "After every draft, plain code checks that each number and date appears in the tool results or in the user's "
        "message, that every citation was retrieved in this turn, that a policy answer using confident results "
        "cites its source, that no other employee's name appears, and that sensitive replies stay within four "
        "sentences. A failed draft is rewritten with the exact problems named, up to twice, and then replaced by a "
        "plain answer built only from the tool output. In the evaluation, 9 of 120 development answers needed a "
        "rewrite, and none of the 180 answers across all three sets fell back to the template. One known gap: "
        "numbers written as words, like 'sixty days', are not checked."
    ),
)
node(s, 0.6, 1.9, 1.9, 0.95, "Draft", ["from the model"], "ext")
node(s, 3.1, 1.9, 1.9, 0.95, "Verify", ["checks"], "security")
node(s, 5.6, 1.9, 1.9, 0.95, "Answer shown", None, "ok")
node(s, 3.1, 3.5, 1.9, 0.95, "Rewrite", ["up to 2 times"], "llm")
node(s, 5.6, 3.5, 1.9, 0.95, "Plain fallback", ["from tool output"], "bad")
arrow(s, [(2.5, 2.37), (3.1, 2.37)])
arrow(s, [(5.0, 2.37), (5.6, 2.37)])
label(s, 5.0, 1.9, 0.6, "pass", size=11.5)
arrow(s, [(4.05, 2.85), (4.05, 3.5)], color=RED)
label(s, 4.15, 3.0, 1.2, "problems", size=11.5, color=RED, align=PP_ALIGN.LEFT)
arrow(s, [(5.0, 3.97), (5.6, 3.97)], dash=True)
label(s, 5.0, 3.6, 0.6, "still", size=11.5)
add_text(s, 0.6, 4.65, 7.0, 0.6, ["A rewrite is checked again. The fallback is built only from tool output and also passes the verifier."], size=13, color=MUTED)
card(s, 8.0, 1.8, 4.75, 3.45, "What it checks", [
    "Every number and date is in the tool results or the user's message",
    "Every citation was retrieved this turn",
    "Policy answers must cite their source",
    "No other employee's name appears",
    "Sensitive replies stay within four sentences",
], accent=AMBER, size=14)
stats = [
    ("9 of 120", "development answers needed a rewrite (7.5%)", BLUE),
    ("0 of 180", "answers fell back to the plain template, across all three sets", TEAL),
    ("1 known gap", "numbers written as words, like “sixty days”, are not checked", RED),
]
for i, (big, small, accent) in enumerate(stats):
    x = 0.6 + i * 4.1
    add_box(s, x, 5.5, 3.95, 1.3, fill=WHITE, line=accent, radius=0.06, line_w=1.75, align=PP_ALIGN.LEFT, margin=0.2,
            paras=[{"runs": [big], "size": 28, "bold": True, "color": accent, "space_after": 0}, {"runs": [small], "size": 13}])

# =====================================================================
# 14  Evaluation
# =====================================================================
s = content_slide(
    "How we know it works",
    "Run against the live system: Claude Sonnet 5.5 answering, Haiku 4.5 routing",
    notes=(
        "Three sets. The held-out set of 30 was written before any results, never tuned against, and run once: 29 of "
        "30, with a 95 percent interval of 83 to 99. The development set of 120 scored 119, but that is the set we "
        "tuned against, so treat it as optimistic. The red-team set of 30 attacks passed 30 of 30 on both runs: "
        "nothing leaked and nothing was written. Almost everything is checked by code rather than by a model judge: "
        "numbers computed from the database by an oracle, citations, which tools were called, and that no unconfirmed "
        "write happened. The evaluation also found a real defect, a sentence counter that merged two sentences after "
        "an email address; it was fixed with a regression test. Say plainly what the numbers cannot show: the same "
        "team wrote the system and the tests, the samples are small, and the world is ten fictional employees."
    ),
)
big = [
    ("29 / 30", "Held-out questions (96.7%)", "95% interval 83% to 99%. Written first, run once, not tuned against.", TEAL),
    ("119 / 120", "Development set (99.2%)", "Optimistic: this is the set used while tuning.", BLUE),
    ("30 / 30", "Red-team attacks (100%)", "Nothing leaked, nothing written. 95% interval 89% to 100%.", PURPLE),
]
for i, (num, head, body, accent) in enumerate(big):
    x = 0.6 + i * 4.1
    add_box(s, x, 1.85, 3.95, 2.1, fill=WHITE, line=accent, radius=0.05, line_w=2, align=PP_ALIGN.LEFT, margin=0.25, anchor=MSO_ANCHOR.TOP,
            paras=[{"runs": [num], "size": 40, "bold": True, "color": accent, "space_after": 0},
                   {"runs": [head], "size": 16, "bold": True, "color": NAVY, "space_after": 4},
                   {"runs": [body], "size": 13}])
add_text(s, 0.6, 4.2, 12.1, 0.4, ["Checked by code, not by a model judge"], size=17, bold=True, color=NAVY)
add_text(s, 0.6, 4.65, 6.0, 1.3, bullets([
    "Numbers match the database oracle: 58/58 development, 12/12 held-out",
    "Cited sections were really retrieved: 58/58 and 14/14",
], space_after=6), size=14.5)
add_text(s, 6.8, 4.65, 6.0, 1.3, bullets([
    "No other person's data in answers: 26/26 and 7/7",
    "No unconfirmed write: 59/59 and 13/13 safety checks",
], space_after=6), size=14.5)
add_box(s, 0.6, 6.1, 12.15, 0.8, fill=AMBER_SOFT, line=None, size=13.5, align=PP_ALIGN.LEFT, margin=0.2,
        paras=["Read with care: the same team wrote the system and the tests, the samples are small, and the world is ten "
               "fictional employees. An independent review is on the roadmap."])

# =====================================================================
# 15  Speed and engineering quality
# =====================================================================
s = content_slide(
    "Speed, and engineering quality",
    notes=(
        "Latency was measured one request at a time. Ordinary questions finish in a median of one to six seconds, "
        "but the complete answer arrives only after verification. Planning takes a median of 32 seconds and 51 "
        "at the 95th percentile, which misses our 30-second target; we say so plainly and it is first on the "
        "improvement list. On quality: 451 automated tests, all passing, a clean linter, real-browser tests of the UI "
        "including phone width, a CI pipeline that has passed on GitHub, a Docker image with a smoke test, and pinned "
        "dependencies. The security policy forbids inline scripts and styles, and every tool call is audited."
    ),
)
add_text(s, 0.6, 1.7, 6.5, 0.4, ["Median seconds to a complete, verified answer"], size=15, bold=True, color=NAVY)
lat = [("Out of scope", 1.0), ("Personal data", 4.1), ("Manager", 4.7), ("Calculation", 5.5), ("Policy", 5.7),
       ("Leave workflow", 6.4), ("Planning", 32.2)]
x0, scale = 2.45, 0.115
for i, (name, val) in enumerate(lat):
    y = 2.25 + i * 0.6
    add_text(s, 0.6, y, 1.8, 0.4, [name], size=14, align=PP_ALIGN.RIGHT)
    colour = PURPLE if name == "Planning" else BLUE
    add_box(s, x0, y + 0.04, max(val * scale, 0.08), 0.34, fill=colour, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, x0 + val * scale + 0.08, y, 1.0, 0.4, [f"{val:g} s"], size=13, bold=True, color=colour)
arrow(s, [(x0 + 30 * scale, 2.2), (x0 + 30 * scale, 6.45)], color=RED, dash=True, head=False, width=1.5)
label(s, x0 + 30 * scale - 0.75, 6.45, 1.5, "30 s target", size=12, color=RED)
add_text(s, 0.6, 6.55, 6.5, 0.45, ["Planning: median 32 s, 95th percentile 51 s. The target is missed."], size=12.5, color=MUTED)
card(s, 7.5, 1.75, 5.25, 5.0, "Engineering quality", [
    "451 automated tests, all passing; linter clean",
    "Real-browser tests of the UI, including phone width",
    "CI on GitHub: first run passed",
    "Docker image with a smoke test; pinned dependency lock",
    "Strict content-security policy: no inline scripts or styles",
    "An audit log row for every tool call, even refused ones",
], accent=TEAL, size=15)

# =====================================================================
# 16  Limits and roadmap
# =====================================================================
s = content_slide(
    "What we know is weak, and what comes next",
    notes=(
        "Close with candour. Known limits: planning is slow; the testing used ten fictional employees and about 50 "
        "sample policy sections; the evaluation was written by the same team; numbers written as words are not "
        "verified; there is no load test and the rate limits live in memory; observability is a trace id and an audit "
        "log, with no tracing yet. Next steps: prompt caching and fewer model round trips for planning, retries with "
        "backoff, tracing, cancelling a request through chat, longer memory, real policies with single sign-on and "
        "HTTPS, and an independent review of the evaluation. One dated item: Haiku 4.5 is announced to retire no "
        "sooner than 15 October 2026, so the router moves to Sonnet before then. Finish on the principle: built to "
        "be trusted first and fast second."
    ),
)
card(s, 0.6, 1.8, 6.0, 4.35, "Known limits", [
    "Planning latency: median 32 s against a 30 s target",
    "Tested on ten fictional employees and about 50 sample policy sections",
    "Evaluation written by the same team; independent review pending",
    "Numbers written as words are not verified",
    "No load test; rate limits are in memory, per process",
    "Observability is a trace id and an audit log; no tracing yet",
], accent=RED, size=14.5)
card(s, 6.8, 1.8, 5.95, 4.35, "Next steps", [
    "Prompt caching and fewer model round trips for planning",
    "Retry with backoff on model calls",
    "Tracing and metrics",
    "Cancel a request through chat; longer-term memory",
    "Real policies, single sign-on and HTTPS",
    "Router moves to Sonnet before Haiku 4.5 retires (no sooner than 2026-10-15)",
], accent=TEAL, size=14.5)
add_box(s, 0.6, 6.3, 12.15, 0.62, fill=NAVY, line=None, paras=["Built to be trusted first, fast second."], size=22, color=WHITE, bold=True)

# =====================================================================
# Appendix
# =====================================================================
s = content_slide(
    "Appendix A: full architecture by layer",
    notes="Reference slide. The same system in nine layers, drawn from docs/ARCHITECTURE.md. Zoom in on the PNG in docs/diagrams/architecture/png for detail.",
)
picture(s, PNG / "architecture-full-zoned.png", 0.6, 1.5, 12.15, 5.4)

s = content_slide(
    "Appendix B: the Deep Agents planning path",
    notes="Reference slide. The planner delegates to two read-only sub-agents inside shared budgets, and falls back to the normal tool loop if it fails.",
)
picture(s, PNG / "03-deep-agents-path.png", 0.6, 1.5, 12.15, 5.4)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"saved {OUT} ({count['n']} slides)")
